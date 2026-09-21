# -*- coding: utf-8 -*-
"""Conversor Markdown -> DOCX sob medida para POPs (tutoriais ilustrados).

Uso:  python md2docx.py entrada.md saida.docx
Dep.: pip install python-docx
Suporta: H1 capa, H2/H3 de marca, imagens centralizadas (retrato/paisagem),
         tabelas, blockquotes como caixas de aviso, listas, negrito/codigo/links,
         blocos de codigo ``` (fonte mono, fundo cinza) e imagens AUTOCONTIDAS
         (data:image/...;base64,...) alem de caminhos relativos.
"""
import base64
import io
import os
import re
import struct
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BRAND = RGBColor(0x1E, 0x40, 0x6B)      # azul escuro p/ títulos
ACCENT = RGBColor(0x25, 0x63, 0xEB)     # azul de destaque
NOTE_BG = "EAF2FB"                       # fundo das caixas de aviso
NOTE_BORDER = "2563EB"
CODE_BG = "F4F4F6"                       # fundo dos blocos de codigo
CODE_FG = RGBColor(0x1B, 0x2B, 0x4A)
BRAND_NAME = os.environ.get("POP_BRAND", "LIS Checklist")  # subtítulo da capa


def png_size_bytes(raw):
    if len(raw) >= 24 and raw[12:16] == b"IHDR":
        w, h = struct.unpack(">II", raw[16:24])
        return w, h
    return 1000, 1000


def png_size(path):
    with open(path, "rb") as f:
        return png_size_bytes(f.read(24))


def resolve_image(rel, base):
    """Retorna (fonte_para_add_picture, w, h) ou None. Aceita data URI ou caminho."""
    if rel.startswith("data:image"):
        try:
            _, b64 = rel.split(",", 1)
            raw = base64.b64decode(b64)
        except Exception:
            return None
        return io.BytesIO(raw), *png_size_bytes(raw[:24])
    img = os.path.normpath(os.path.join(base, rel))
    if os.path.exists(img):
        return img, *png_size(img)
    return None


def shade_paragraph(p, fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def left_border(p, color, sz=24):
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    el = OxmlElement("w:left")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), "8")
    el.set(qn("w:color"), color)
    pbdr.append(el)
    pPr.append(pbdr)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def add_runs(paragraph, text, base_bold=False):
    """Adiciona texto com **negrito**, `código` e [link](url)->texto."""
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = paragraph.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(0xB0, 0x2A, 0x37)
        elif part.startswith("[") and "](" in part:
            label = part[1:part.index("]")]
            r = paragraph.add_run(label)
            r.font.color.rgb = ACCENT
        else:
            r = paragraph.add_run(part)
        if base_bold:
            r.bold = True


def add_code_block(doc, code_lines):
    """Bloco de codigo: fonte mono, fundo cinza, borda esquerda."""
    for cl in code_lines:
        p = doc.add_paragraph()
        shade_paragraph(p, CODE_BG)
        left_border(p, "C8CCD4", sz=18)
        pf = p.paragraph_format
        pf.left_indent = Pt(8)
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        r = p.add_run(cl if cl else " ")
        r.font.name = "Consolas"
        r.font.size = Pt(9.5)
        r.font.color.rgb = CODE_FG
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def convert(md_path, docx_path):
    base = os.path.dirname(os.path.abspath(md_path))
    with open(md_path, encoding="utf-8") as f:
        lines = f.read().split("\n")

    doc = Document()

    # estilo base
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    for sec in doc.sections:
        sec.top_margin = Inches(0.9)
        sec.bottom_margin = Inches(0.9)
        sec.left_margin = Inches(1.0)
        sec.right_margin = Inches(1.0)

    i = 0
    first_h1 = True
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # bloco de codigo cercado ```
        if stripped.startswith("```"):
            i += 1
            code = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1  # pula a cerca de fechamento
            add_code_block(doc, code)
            continue

        # vazio
        if not stripped:
            i += 1
            continue

        # imagem (caminho relativo OU data URI autocontida)
        m = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if m:
            res = resolve_image(m.group(2), base)
            if res:
                src, w, h = res
                width = Inches(6.4) if w >= h else Inches(2.95)
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(src, width=width)
                p.space_after = Pt(6)
            i += 1
            continue

        # cabeçalhos
        if stripped.startswith("### "):
            p = doc.add_heading(level=2)
            add_runs(p, stripped[4:])
            for r in p.runs:
                r.font.color.rgb = BRAND
            i += 1
            continue
        if stripped.startswith("## "):
            p = doc.add_heading(level=1)
            add_runs(p, stripped[3:])
            for r in p.runs:
                r.font.color.rgb = BRAND
            i += 1
            continue
        if stripped.startswith("# "):
            txt = stripped[2:]
            if first_h1:
                first_h1 = False
                t = doc.add_paragraph()
                t.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = t.add_run(txt)
                run.bold = True
                run.font.size = Pt(26)
                run.font.color.rgb = BRAND
                sub = doc.add_paragraph()
                sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
                sr = sub.add_run(BRAND_NAME)
                sr.font.size = Pt(13)
                sr.font.color.rgb = ACCENT
            else:
                p = doc.add_heading(level=0)
                add_runs(p, txt)
            i += 1
            continue

        # regra horizontal
        if re.match(r"^-{3,}$", stripped):
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pbdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), "C8C8C8")
            pbdr.append(bottom)
            pPr.append(pbdr)
            i += 1
            continue

        # blockquote / aviso
        if stripped.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(re.sub(r"^\s*>\s?", "", lines[i]))
                i += 1
            text = " ".join(x.strip() for x in buf if x.strip())
            p = doc.add_paragraph()
            shade_paragraph(p, NOTE_BG)
            left_border(p, NOTE_BORDER)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.left_indent = Pt(8)
            add_runs(p, text)
            continue

        # tabela
        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip())
                i += 1
            parsed = []
            for r in rows:
                if re.match(r"^\|[\s:|-]+\|$", r):  # separador
                    continue
                cells = [c.strip() for c in r.strip().strip("|").split("|")]
                parsed.append(cells)
            if parsed:
                ncol = len(parsed[0])
                table = doc.add_table(rows=0, cols=ncol)
                table.style = "Light Grid Accent 1"
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                for ri, cells in enumerate(parsed):
                    cs = table.add_row().cells
                    for ci in range(ncol):
                        cell = cs[ci]
                        cell.text = ""
                        para = cell.paragraphs[0]
                        add_runs(para, cells[ci] if ci < len(cells) else "", base_bold=(ri == 0))
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        # lista ordenada — numeração manual (preserva o número do Markdown)
        m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.35)
            p.paragraph_format.first_line_indent = Inches(-0.35)
            p.paragraph_format.space_after = Pt(2)
            num_run = p.add_run(f"{m.group(1)}.  ")
            num_run.bold = True
            add_runs(p, m.group(2))
            i += 1
            continue
        # lista não ordenada
        if stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, stripped[2:])
            i += 1
            continue

        # parágrafo normal
        p = doc.add_paragraph()
        add_runs(p, stripped)
        i += 1

    doc.save(docx_path)
    print("OK ->", docx_path)


if __name__ == "__main__":
    convert(sys.argv[1], sys.argv[2])
